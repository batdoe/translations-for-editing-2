"""Shared helpers for the Inn Trouble proofread tools (cloud stand-ins for the PC scripts)."""
import os, re, glob
from docx import Document
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CORRECTED = os.path.join(ROOT, "corrected")
BASELINE = os.path.join(CORRECTED, "_baseline")   # post-sweep snapshot the verify step compares against
LEDGERS = os.path.join(ROOT, "_EditLedgers")
ENGLISH = os.path.join(ROOT, "_source", "Inn Trouble - ASSEMBLED - FORMATTED.docx")

SCENE_BREAK = re.compile(r"^\s*(\*\s*){3}\s*$")
BT = re.compile(r"<bt>.*?</bt>")


def stem_path(stem, where="translated"):
    base = {"translated": ROOT, "corrected": CORRECTED, "baseline": BASELINE}[where]
    return os.path.join(base, stem + ".docx")


def editions():
    return sorted(os.path.splitext(os.path.basename(f))[0] for f in glob.glob(os.path.join(ROOT, "*_*.docx")))


def is_italic(run_el):
    rpr = run_el.find(qn("w:rPr"))
    if rpr is None:
        return False
    i = rpr.find(qn("w:i"))
    if i is None:
        return False
    val = i.get(qn("w:val"))
    return val not in ("0", "false", "off")


def italic_runs(doc):
    """Text of every italic run, in document order (parsed, never string-counted)."""
    out = []
    for p in doc.paragraphs:
        for r in p._p.iter(qn("w:r")):
            if is_italic(r):
                t = "".join(x.text or "" for x in r.iter(qn("w:t")))
                if t.strip():
                    out.append(t)
    return out


def stats(path):
    d = Document(path)
    ps = d.paragraphs
    text = [p.text for p in ps]
    return {
        "paras": len(ps),
        "nonempty": sum(1 for t in text if t.strip()),
        "heading1": sum(1 for p in ps if p.style.name == "Heading 1"),
        "scene_breaks": sum(1 for t in text if SCENE_BREAK.match(t)),
        "asterisks": sum(t.count("*") for t in text if not SCENE_BREAK.match(t)),
        "bt_spans": sum(len(BT.findall(t)) for t in text),
        "italics": italic_runs(d),
        "text": text,
    }
